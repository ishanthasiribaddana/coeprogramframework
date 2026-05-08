from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.util import Pt
from pptx.enum.text import PP_ALIGN
from pptx.dml.color import RGBColor

# Create presentation
prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)

def add_title_slide(title, subtitle):
    slide_layout = prs.slide_layouts[6]  # Blank
    slide = prs.slides.add_slide(slide_layout)
    
    # Title
    title_box = slide.shapes.add_textbox(Inches(0.5), Inches(2.5), Inches(12.333), Inches(1.5))
    tf = title_box.text_frame
    p = tf.paragraphs[0]
    p.text = title
    p.font.size = Pt(54)
    p.font.bold = True
    p.font.color.rgb = RGBColor(0x1E, 0x3A, 0x8A)
    p.alignment = PP_ALIGN.CENTER
    
    # Subtitle
    sub_box = slide.shapes.add_textbox(Inches(0.5), Inches(4.2), Inches(12.333), Inches(1))
    tf = sub_box.text_frame
    p = tf.paragraphs[0]
    p.text = subtitle
    p.font.size = Pt(28)
    p.font.color.rgb = RGBColor(0x64, 0x74, 0x8B)
    p.alignment = PP_ALIGN.CENTER

def add_content_slide(title, bullets):
    slide_layout = prs.slide_layouts[6]  # Blank
    slide = prs.slides.add_slide(slide_layout)
    
    # Title
    title_box = slide.shapes.add_textbox(Inches(0.5), Inches(0.4), Inches(12.333), Inches(1))
    tf = title_box.text_frame
    p = tf.paragraphs[0]
    p.text = title
    p.font.size = Pt(40)
    p.font.bold = True
    p.font.color.rgb = RGBColor(0x1E, 0x3A, 0x8A)
    
    # Content
    content_box = slide.shapes.add_textbox(Inches(0.7), Inches(1.5), Inches(12), Inches(5.5))
    tf = content_box.text_frame
    tf.word_wrap = True
    
    for i, bullet in enumerate(bullets):
        if i == 0:
            p = tf.paragraphs[0]
        else:
            p = tf.add_paragraph()
        p.text = bullet
        p.font.size = Pt(24)
        p.font.color.rgb = RGBColor(0x33, 0x33, 0x33)
        p.space_before = Pt(12)
        p.level = 0

def add_two_column_slide(title, left_content, right_content):
    slide_layout = prs.slide_layouts[6]
    slide = prs.slides.add_slide(slide_layout)
    
    # Title
    title_box = slide.shapes.add_textbox(Inches(0.5), Inches(0.4), Inches(12.333), Inches(1))
    tf = title_box.text_frame
    p = tf.paragraphs[0]
    p.text = title
    p.font.size = Pt(40)
    p.font.bold = True
    p.font.color.rgb = RGBColor(0x1E, 0x3A, 0x8A)
    
    # Left column
    left_box = slide.shapes.add_textbox(Inches(0.5), Inches(1.5), Inches(6), Inches(5.5))
    tf = left_box.text_frame
    tf.word_wrap = True
    for i, item in enumerate(left_content):
        if i == 0:
            p = tf.paragraphs[0]
        else:
            p = tf.add_paragraph()
        p.text = item
        p.font.size = Pt(22)
        p.space_before = Pt(8)
    
    # Right column
    right_box = slide.shapes.add_textbox(Inches(6.8), Inches(1.5), Inches(6), Inches(5.5))
    tf = right_box.text_frame
    tf.word_wrap = True
    for i, item in enumerate(right_content):
        if i == 0:
            p = tf.paragraphs[0]
        else:
            p = tf.add_paragraph()
        p.text = item
        p.font.size = Pt(22)
        p.space_before = Pt(8)

# Slide 1: Title
add_title_slide(
    "COE Program Framework",
    "A Digital Platform for Program Management\nAnanda College Center of Excellence"
)

# Slide 2: Introduction
add_content_slide("What is COE Program Framework?", [
    "• A comprehensive web application for managing educational programs",
    "• Designed for Ananda College Center of Excellence",
    "• Enables 10 specialized centers to plan, track, and report programs",
    "• Streamlines program identification and documentation",
    "• Real-time collaboration and data synchronization"
])

# Slide 3: The Challenge
add_content_slide("The Challenge - Before the Framework", [
    "• Manual program tracking across multiple centers",
    "• Inconsistent documentation formats",
    "• Difficulty in cross-center collaboration",
    "• Time-consuming report generation",
    "• No centralized data repository",
    "• Limited visibility across departments"
])

# Slide 4: The Solution
add_content_slide("The Solution - COE Program Framework", [
    "✓ Centralized Platform - Single source of truth for all programs",
    "✓ Real-time Updates - Instant synchronization across users",
    "✓ Automated Reports - One-click PDF generation",
    "✓ Cross-Center Integration - Seamless collaboration",
    "✓ Cloud-based Storage - Secure and accessible anywhere"
])

# Slide 5: Centers Overview
add_two_column_slide("10 Specialized Centers", [
    "🤖 AI Center",
    "⚡ STEAM Hub",
    "📚 Language Center",
    "🎭 Auditorium",
    "🔬 Science Center"
], [
    "💼 Entrepreneurship Center",
    "🎬 Media Center",
    "📐 Mathematics Center",
    "🎨 Fine Arts Center",
    "🎵 Performing Arts Center"
])

# Slide 6: Program Types
add_content_slide("Three Categories of Programs", [
    "1. ADVANCED PROGRAMS",
    "   • Employment-ready skills not in current curriculum",
    "   • Core programs specific to each center",
    "",
    "2. STEAM PROGRAMS",
    "   • Cross-disciplinary integration",
    "   • Science, Technology, Engineering, Arts, Mathematics",
    "",
    "3. CROSS-CENTER PROGRAMS",
    "   • Inter-center collaboration",
    "   • Shared resources and expertise"
])

# Slide 7: Key Features
add_content_slide("Application Features", [
    "• Center Selection - Visual icons for easy navigation",
    "• Program Management - Add, edit, remove programs",
    "• Partnership Tracking - External and placement partners",
    "• Student Associations - Link programs to student groups",
    "• Report Generation - Professional PDF exports",
    "• Auto-save - Real-time data persistence"
])

# Slide 8: Program Details
add_content_slide("Information Captured Per Program", [
    "• Module Name - Program/project title",
    "• Duration - Estimated hours (e.g., 20-30 hrs)",
    "• External Partnerships - Industry collaborators",
    "• Career Guidance - Placement partners",
    "• Student Associations - Linked student groups",
    "• Cross Center - Partnering center (if applicable)"
])

# Slide 9: Student Associations
add_content_slide("Linked Student Organizations", [
    "• Information and Communication Technology Unit",
    "• Inventors' Association",
    "• Entrepreneurship Association",
    "• Engineering Technology Association",
    "• Robotics Association",
    "• Green Energy Association",
    "• Research and Exploration Society"
])

# Slide 10: Technology Stack
add_two_column_slide("Modern Technology Stack", [
    "FRONTEND:",
    "• React 18 - User Interface",
    "• Tailwind CSS - Styling",
    "• Vite - Build Tool",
    "• Lucide Icons"
], [
    "BACKEND:",
    "• Node.js - Server Runtime",
    "• Express.js - API Framework",
    "• MySQL/MariaDB - Database",
    "• JWT - Authentication"
])

# Slide 11: Architecture
add_content_slide("System Architecture", [
    "┌─────────────┐    ┌─────────────┐    ┌─────────────┐",
    "│ Web Browser │ → │ Nginx/SSL   │ → │ Node.js API │",
    "│ (React App) │    │ (Port 443)  │    │ (Port 3004) │",
    "└─────────────┘    └─────────────┘    └──────┬──────┘",
    "                                              ↓",
    "                                    ┌─────────────────┐",
    "                                    │ MySQL Database  │",
    "                                    └─────────────────┘"
])

# Slide 12: Database Design
add_content_slide("Normalized Database Schema (3NF)", [
    "CORE TABLES:",
    "• Centers, Programs, Program Types",
    "• External Partners, Placement Partners, Student Associations",
    "",
    "JUNCTION TABLES:",
    "• Program-Partners, Program-Placements, Program-Associations",
    "",
    "AUDIT TABLES:",
    "• Submissions, Submission Notes, Audit Log"
])

# Slide 13: Security
add_content_slide("Security Features", [
    "🔒 HTTPS/SSL - Encrypted data transmission",
    "🔑 JWT Authentication - Secure user sessions",
    "👥 Role-based Access - Center-specific permissions",
    "📝 Audit Logging - Track all changes",
    "💾 Database Backup - Regular data backups",
    "🛡️ Input Validation - Protection against attacks"
])

# Slide 14: Current Status
add_content_slide("Current Implementation Status", [
    "✅ AI Center - 15+ programs finalized",
    "✅ STEAM Hub - 10+ programs finalized",
    "📝 Other Centers - Ready for input",
    "",
    "TOTAL: 43+ programs imported from master schedule",
    "",
    "Live URL: https://coe-sedf.oapf.org/"
])

# Slide 15: Benefits
add_content_slide("Benefits Delivered", [
    "✅ EFFICIENCY - Reduced documentation time by 70%",
    "✅ CONSISTENCY - Standardized program formats",
    "✅ COLLABORATION - Easy cross-center coordination",
    "✅ ACCESSIBILITY - Available 24/7 from any device",
    "✅ SCALABILITY - Ready for future expansion",
    "✅ TRANSPARENCY - Real-time visibility for all stakeholders"
])

# Slide 16: Future Roadmap
add_content_slide("Future Enhancements", [
    "• Dashboard Analytics - Visual program statistics",
    "• Mobile App - Native iOS/Android support",
    "• Approval Workflow - Multi-level program approval",
    "• Calendar Integration - Schedule visualization",
    "• Email Notifications - Automated alerts",
    "• Export to Excel - Bulk data export"
])

# Slide 17: Thank You
add_title_slide(
    "Thank You",
    "COE Program Framework\nEmpowering Excellence Through Technology\n\nhttps://coe-sedf.oapf.org/"
)

# Save presentation
output_path = r'd:\Winsurf\OAPF\coe-program-framework\COE_Program_Framework.pptx'
prs.save(output_path)
print(f"Presentation saved to: {output_path}")
